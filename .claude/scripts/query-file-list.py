#!/usr/bin/env python3
"""
File List Query — 从 file-list.xlsx 按条件查询文件列表（generate-file-list.py 产出）。

用法: python3 query-file-list.py <project_root> [选项]

选项:
  --type <types>            按文件类型过滤，逗号分隔。如: Entity,Controller
                            合法值动态校验，非法值 fail-fast；运行 --list-types 查看完整列表
  --list-types              打印所有合法 fileType 后退出（无需 project_root）
  --package <pattern>       按包名过滤，子串匹配。如: service.campaign
  --exclude-package <pkg>   排除包含此子串的包名。可多次指定
  --name <pattern>          按文件名过滤，支持正则。如: Campaign.*|Campaign
  --exclude-name <pattern>  排除匹配此正则的文件名。可多次指定
  --exclude-demo           排除包路径或类名命中 demo/example/sample/test 的记录（示例代码不参与建域）
  --group-by <field>        按字段分组输出: type / package。默认不分组
  --limit <N>               限制输出条数，默认不限
  --json                    以 JSON 格式输出

示例:
  python3 query-file-list.py . --type Entity
  python3 query-file-list.py . --type Controller,ServiceInterface --package campaign
  python3 query-file-list.py . --type Entity,Controller --group-by package
  python3 query-file-list.py . --name "Campaign.*|Advert.*"
  python3 query-file-list.py . --type Entity --json
  python3 query-file-list.py . --type Entity --exclude-package campaign --exclude-package adunit --json
"""

import json
import os
import re
import sys
from collections import defaultdict

try:
    from openpyxl import load_workbook
except ImportError:
    print("错误: 需要安装 openpyxl — pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)

# ── fileType 词典（单一信源） ──
# 须与 generate-file-list.py 的 TYPE_ORDER 保持一致；本常量为对外契约，调用方/校验脚本以此为准。
FILE_TYPES = (
    "Entity", "Controller", "ServiceInterface", "ServiceImpl",
    "MapperJava", "MapperXml", "Consumer", "Job", "Dto",
    "Test", "Util", "Other",
)
FILE_TYPE_SET = frozenset(FILE_TYPES)

# demo/example/sample/test 判定 — 口径与 generate-file-list.py 保持一致（包路径段 OR 类名驼峰词）
DEMO_SEGMENTS = {"demo", "example", "sample", "test"}


def is_demo_record(rec: dict) -> bool:
    stem = rec["fileName"].rsplit(".", 1)[0] if "." in rec["fileName"] else rec["fileName"]
    segs = {s.lower() for s in rec["package"].split(".") if s}
    words = {w.lower() for w in re.findall(r"[A-Z]+(?![a-z])|[A-Z][a-z0-9]*", stem)}
    return bool((segs | words) & DEMO_SEGMENTS)


def split_fqn(fqn: str) -> tuple[str, str]:
    """FQN → (package, simpleClassName)。无包名时返回 ("", fqn)。"""
    if "." not in fqn:
        return "", fqn
    pkg, _, simple = fqn.rpartition(".")
    return pkg, simple


def load_file_list(project_root: str) -> list[dict]:
    """从 file-list.xlsx 读取，返回扁平记录列表。"""
    path = os.path.join(project_root, "knowledge", "file-list.xlsx")
    if not os.path.exists(path):
        print(f"错误: {path} 不存在，请先运行 generate-file-list.py", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    records = []
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not header_seen:
            header_seen = True  # 跳过表头
            continue
        if not row or not row[0]:
            continue
        fqn = str(row[0]).strip()
        file_suffix = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        file_type = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        pkg, simple_name = split_fqn(fqn)
        records.append({
            "fileName": f"{simple_name}.{file_suffix}" if file_suffix else simple_name,
            "fileType": file_type,
            "package": pkg,
        })
    wb.close()
    return records


def query(records: list[dict], types: list[str] | None, package: str | None,
          name_pattern: str | None, exclude_packages: list[str] | None = None,
          exclude_name_patterns: list[str] | None = None,
          exclude_demo: bool = False) -> list[dict]:
    results = []
    name_re = re.compile(name_pattern, re.IGNORECASE) if name_pattern else None
    exclude_name_res = [re.compile(p, re.IGNORECASE) for p in exclude_name_patterns] if exclude_name_patterns else []
    for r in records:
        if types and r["fileType"] not in types:
            continue
        if exclude_demo and is_demo_record(r):
            continue
        if package and package not in r["package"]:
            continue
        if name_re and not name_re.search(r["fileName"]):
            continue
        if exclude_packages and any(ep in r["package"] for ep in exclude_packages):
            continue
        if exclude_name_res and any(er.search(r["fileName"]) for er in exclude_name_res):
            continue
        results.append(r)
    return results


def print_table(records: list[dict], group_by: str | None):
    if group_by == "type":
        groups = defaultdict(list)
        for r in records:
            groups[r["fileType"]].append(r)
        for ftype in sorted(groups.keys()):
            items = groups[ftype]
            print(f"\n--- {ftype} ({len(items)}) ---")
            for r in sorted(items, key=lambda x: (x["package"], x["fileName"])):
                print(f"  {r['package']}.{r['fileName']}" if r['package'] else f"  {r['fileName']}")
    elif group_by == "package":
        groups = defaultdict(list)
        for r in records:
            groups[r["package"]].append(r)
        for pkg in sorted(groups.keys()):
            items = groups[pkg]
            types = defaultdict(list)
            for r in items:
                types[r["fileType"]].append(r["fileName"])
            print(f"\n--- {pkg or '(无包名)'} ({len(items)}) ---")
            for ftype in sorted(types.keys()):
                names = sorted(types[ftype])
                print(f"  {ftype}: {names}")
    else:
        for r in sorted(records, key=lambda x: (x["fileType"], x["package"], x["fileName"])):
            qualified = f"{r['package']}.{r['fileName']}" if r['package'] else r['fileName']
            print(f"  {r['fileType']:<18} {qualified}")


def main():
    # --list-types：打印合法 fileType 后退出（无需 project_root）
    if "--list-types" in sys.argv[1:]:
        print("\n".join(FILE_TYPES))
        sys.exit(0)

    if len(sys.argv) < 2:
        print(__doc__, file=sys.stderr)
        sys.exit(1)

    project_root = os.path.abspath(sys.argv[1])

    types = None
    package = None
    name_pattern = None
    exclude_packages = []
    exclude_name_patterns = []
    exclude_demo = False
    group_by = None
    limit = None
    output_json = False

    args = sys.argv[2:]
    i = 0
    while i < len(args):
        if args[i] == "--type" and i + 1 < len(args):
            types = [t.strip() for t in args[i + 1].split(",")]
            i += 2
        elif args[i] == "--package" and i + 1 < len(args):
            package = args[i + 1]
            i += 2
        elif args[i] == "--exclude-package" and i + 1 < len(args):
            exclude_packages.append(args[i + 1])
            i += 2
        elif args[i] == "--exclude-name" and i + 1 < len(args):
            exclude_name_patterns.append(args[i + 1])
            i += 2
        elif args[i] == "--name" and i + 1 < len(args):
            name_pattern = args[i + 1]
            i += 2
        elif args[i] == "--group-by" and i + 1 < len(args):
            group_by = args[i + 1]
            i += 2
        elif args[i] == "--limit" and i + 1 < len(args):
            limit = int(args[i + 1])
            i += 2
        elif args[i] == "--exclude-demo":
            exclude_demo = True
            i += 1
        elif args[i] == "--json":
            output_json = True
            i += 1
        else:
            print(f"未知参数: {args[i]}", file=sys.stderr)
            sys.exit(1)

    # fileType 合法性校验（fail-fast，统一词典，拦截 Mapper-JAVA 之类拼写错误）
    if types:
        invalid = [t for t in types if t not in FILE_TYPE_SET]
        if invalid:
            print(
                f"错误: 非法 fileType: {','.join(invalid)}\n"
                f"合法值: {','.join(FILE_TYPES)}\n"
                f"提示: 运行 python3 query-file-list.py --list-types 查看完整列表",
                file=sys.stderr,
            )
            sys.exit(1)

    records = load_file_list(project_root)
    results = query(records, types, package, name_pattern, exclude_packages, exclude_name_patterns, exclude_demo)
    if limit:
        results = results[:limit]

    if output_json:
        print(json.dumps(results, ensure_ascii=False, indent=2))
    else:
        print(f"查询结果: {len(results)} 条" + (f" (限制 {limit})" if limit else ""))
        print("-" * 60)
        print_table(results, group_by)
        print("-" * 60)
        print(f"合计: {len(results)} 条")


if __name__ == "__main__":
    main()

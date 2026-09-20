#!/usr/bin/env python3
"""
Entry Parser — 从 file-list.xlsx 召回 Controller/Job/Consumer 文件，解析入口方法签名，输出结构化 JSON。

用法: python3 parse-entry.py <project_root> [--output <path>]

产出:
  - <output>  结构化 JSON（包含 Controller/Job/Consumer 三类入口的方法列表）
  - stdout    统计摘要

入口类型与解析内容:
  Controller:
    - 类级别 @RequestMapping 路径前缀
    - 方法 HTTP 注解（@GetMapping/@PostMapping/@PutMapping/@DeleteMapping/@PatchMapping/@RequestMapping）
    - 路径、HTTP 方法
  Job:
    - @XxlJob("任务名")
    - @Scheduled(cron="...") / @Scheduled(fixedRate=...)
  Consumer:
    - @KafkaListener(topics="...")
    - @RabbitListener(queues="...")
    - @RocketMQMessageListener(topic="...", consumerGroup="...")（类级注解）
  通用:
    - 方法名、返回类型
    - 参数 DTO/VO 类型
    - 方法注释（JavaDoc 或 // 单行注释）
"""

import json
import os
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("错误: 需要安装 openpyxl — pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)


# HTTP 注解 → HTTP 方法映射
HTTP_ANNOTATIONS = {
    "GetMapping": "GET",
    "PostMapping": "POST",
    "PutMapping": "PUT",
    "DeleteMapping": "DELETE",
    "PatchMapping": "PATCH",
    "RequestMapping": "REQUEST",
}

# file-list.xlsx 中视为入口的 fileType → entry_type
ENTRY_FILE_TYPES = {
    "Controller": "Controller",
    "Job": "Job",
    "Consumer": "Consumer",
}

# 技术参数类型，不参与业务实体推断
TECH_PARAM_TYPES = {
    "HttpServletRequest", "HttpServletResponse", "HttpSession",
    "Model", "ModelMap", "ModelAndView",
    "BindingResult", "Errors",
    "MultipartFile", "MultipartHttpServletRequest",
    "RedirectAttributes", "UriComponentsBuilder",
    "WebRequest", "NativeWebRequest",
    "InputStream", "OutputStream", "Reader", "Writer",
    "Principal", "Authentication",
    "Long", "Integer", "String", "Boolean", "int", "long", "boolean",
    "List", "Map", "Set", "Collection",
    "Pageable", "PageRequest", "Sort",
}

# 通用方法签名尾部：匹配 public/protected 方法声明
# 捕获 return_type、method_name、params
COMMON_METHOD_TAIL = (
    r'(?:@\w+(?:\([^)]*\))?\s*)*'
    r'(?:public|protected)\s+'
    r'(?:static\s+)?'
    r'(?:synchronized\s+)?'
    r'(?P<return_type>\S+(?:\s*<\s*\S+\s*(?:,\s*\S+\s*)*>)?)\s+'
    r'(?P<method_name>\w+)\s*'
    r'\((?P<params>[^)]*)\)'
)


def load_entry_files(project_root: str) -> list[dict]:
    """从 file-list.xlsx 召回 Controller / Job / Consumer 文件"""
    path = os.path.join(project_root, "knowledge", "file-list.xlsx")
    if not os.path.exists(path):
        print(f"错误: {path} 不存在，请先运行 generate-file-list.py", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    entries = []
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not header_seen:
            header_seen = True  # 跳过表头
            continue
        if not row or not row[0]:
            continue
        fqn = str(row[0]).strip()
        file_suffix = str(row[1]).strip() if len(row) > 1 and row[1] else "java"
        file_type = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        if file_type not in ENTRY_FILE_TYPES:
            continue
        if "." in fqn:
            pkg, _, simple = fqn.rpartition(".")
        else:
            pkg, simple = "", fqn
        entries.append({
            "entry_type": ENTRY_FILE_TYPES[file_type],
            "package_name": pkg,
            "file_name": simple,
            "file_suffix": file_suffix,
        })
    wb.close()
    return entries


def find_entry_path(project_root: str, package_name: str, file_name: str) -> str | None:
    """根据包名和文件名搜索入口源文件的实际路径"""
    root = Path(project_root).resolve()
    pkg_path = package_name.replace(".", "/")
    base_name = file_name

    for src_dir in root.rglob("src/main/java"):
        candidate = src_dir / pkg_path / f"{base_name}.java"
        if candidate.exists():
            return str(candidate)

    for f in root.rglob(f"{base_name}.java"):
        try:
            content = f.read_text(encoding="utf-8", errors="ignore")
            if f"package {package_name}" in content:
                return str(f)
        except Exception:
            continue

    return None


def extract_package_prefix(package_name: str) -> str:
    """提取业务包前缀（去掉技术层后缀）"""
    parts = package_name.split(".")
    tech_suffixes = {
        "controller", "service", "impl", "mapper", "dao",
        "entity", "model", "po", "do", "dto", "vo", "param",
        "job", "task", "scheduler", "config", "filter", "interceptor",
        "consumer", "listener", "provider", "handler", "util", "common",
        "enums", "constant", "constants", "exception", "rpc", "feign",
        "api",
    }
    result = list(parts)
    while result and result[-1].lower() in tech_suffixes:
        result.pop()
    return ".".join(result)


def extract_request_mapping(content: str) -> str | None:
    """提取类级别 @RequestMapping 路径"""
    m = re.search(r'@RequestMapping\s*\(\s*(?:value\s*=\s*)?["\']([^"\']+)["\']', content)
    return m.group(1) if m else None


def extract_method_comment(lines: list[str], method_line_idx: int) -> str:
    """提取方法前的 JavaDoc 或单行注释"""
    comments = []
    i = method_line_idx - 1
    while i >= 0:
        line = lines[i].strip()
        if line.startswith("//"):
            comments.insert(0, line.lstrip("/ "))
        elif line.startswith("*") or line.startswith("/**") or line == "*/":
            text = re.sub(r'^\*\s*', '', line).strip()
            if text and not text.startswith("@") and text != "/" and text != "*/":
                comments.insert(0, text)
        elif line == "":
            pass
        else:
            break
        i -= 1
    return " ".join(comments) if comments else ""


def extract_method_params(params_str: str) -> list[str]:
    """从方法参数串中提取业务 DTO/VO 类型"""
    if not params_str or params_str.strip() == "":
        return []

    param_parts = []
    depth = 0
    current = []
    for ch in params_str:
        if ch == "<":
            depth += 1
        elif ch == ">":
            depth -= 1
        elif ch == "," and depth == 0:
            param_parts.append("".join(current).strip())
            current = []
            continue
        current.append(ch)
    if current:
        param_parts.append("".join(current).strip())

    business_types = []
    for param in param_parts:
        param = param.strip()
        param = re.sub(r'@\w+(\([^)]*\))?', '', param).strip()
        parts = param.split()
        if len(parts) >= 1:
            type_name = parts[0]
            simple_type = re.sub(r'.*[.<](\w+)[>]?$', r'\1', type_name)
            if simple_type and simple_type[0].isupper() and simple_type not in TECH_PARAM_TYPES:
                business_types.append(simple_name(type_name))

    return business_types


def simple_name(full_type: str) -> str:
    """从完整类型名提取简单类名"""
    inner = re.search(r'<(\w+)>', full_type)
    if inner:
        return inner.group(1)
    if "." in full_type:
        return full_type.split(".")[-1]
    return full_type


def parse_controller_methods(content: str, lines: list[str]) -> list[dict]:
    """解析 Controller 文件的 HTTP 注解方法"""
    methods = []
    pattern = re.compile(
        r'@(GetMapping|PostMapping|PutMapping|DeleteMapping|PatchMapping|RequestMapping)'
        r'\s*\(\s*(?:value\s*=\s*)?(?:path\s*=\s*)?["\']([^"\']*)["\']'
        r'(?:[^)]*)\)'
        r'\s*'
        + COMMON_METHOD_TAIL,
        re.MULTILINE
    )
    for m in pattern.finditer(content):
        http_anno = m.group(1)
        path = m.group(2) or ""
        return_type = simple_name(m.group("return_type"))
        method_name = m.group("method_name")
        params_str = m.group("params")
        line_idx = content[:m.start()].count("\n")
        comment = extract_method_comment(lines, line_idx)
        param_types = extract_method_params(params_str)
        methods.append({
            "method_name": method_name,
            "entry_annotation": http_anno,
            "entry_descriptor": path,
            "http_method": HTTP_ANNOTATIONS.get(http_anno, "REQUEST"),
            "http_annotation": http_anno,
            "path": path,
            "return_type": return_type,
            "param_types": param_types,
            "comment": comment,
        })
    return methods


def parse_job_methods(content: str, lines: list[str]) -> list[dict]:
    """解析 Job 文件中的 @XxlJob / @Scheduled 方法"""
    methods = []

    xxljob_pattern = re.compile(
        r'@XxlJob\s*\(\s*(?:value\s*=\s*)?["\']([^"\']+)["\']'
        r'(?:[^)]*)\)'
        r'\s*'
        + COMMON_METHOD_TAIL,
        re.MULTILINE
    )
    for m in xxljob_pattern.finditer(content):
        job_name = m.group(1)
        return_type = simple_name(m.group("return_type"))
        method_name = m.group("method_name")
        params_str = m.group("params")
        line_idx = content[:m.start()].count("\n")
        comment = extract_method_comment(lines, line_idx)
        param_types = extract_method_params(params_str)
        methods.append({
            "method_name": method_name,
            "entry_annotation": "XxlJob",
            "entry_descriptor": job_name,
            "return_type": return_type,
            "param_types": param_types,
            "comment": comment,
        })

    scheduled_pattern = re.compile(
        r'@Scheduled\s*\(\s*([^)]+)\)'
        r'\s*'
        + COMMON_METHOD_TAIL,
        re.MULTILINE
    )
    for m in scheduled_pattern.finditer(content):
        attrs = m.group(1)
        return_type = simple_name(m.group("return_type"))
        method_name = m.group("method_name")
        params_str = m.group("params")
        line_idx = content[:m.start()].count("\n")
        comment = extract_method_comment(lines, line_idx)
        param_types = extract_method_params(params_str)
        cron_m = re.search(r'cron\s*=\s*["\']([^"\']+)["\']', attrs)
        if cron_m:
            descriptor = f"cron={cron_m.group(1)}"
        else:
            fr_m = re.search(r'fixed(?:Rate|Delay)\s*=\s*(\d+)', attrs)
            descriptor = fr_m.group(0) if fr_m else attrs.strip()
        methods.append({
            "method_name": method_name,
            "entry_annotation": "Scheduled",
            "entry_descriptor": descriptor,
            "return_type": return_type,
            "param_types": param_types,
            "comment": comment,
        })
    return methods


def parse_consumer_methods(content: str, lines: list[str]) -> list[dict]:
    """解析 Consumer 文件中的 @KafkaListener / @RabbitListener / @RocketMQMessageListener 方法"""
    methods = []

    kafka_pattern = re.compile(
        r'@KafkaListener\s*\(\s*([^)]+)\)'
        r'\s*'
        + COMMON_METHOD_TAIL,
        re.MULTILINE
    )
    for m in kafka_pattern.finditer(content):
        attrs = m.group(1)
        return_type = simple_name(m.group("return_type"))
        method_name = m.group("method_name")
        params_str = m.group("params")
        line_idx = content[:m.start()].count("\n")
        comment = extract_method_comment(lines, line_idx)
        param_types = extract_method_params(params_str)
        topic_m = re.search(r'topics\s*=\s*["\']([^"\']+)["\']', attrs)
        descriptor = f"topic={topic_m.group(1)}" if topic_m else attrs.strip()
        methods.append({
            "method_name": method_name,
            "entry_annotation": "KafkaListener",
            "entry_descriptor": descriptor,
            "return_type": return_type,
            "param_types": param_types,
            "comment": comment,
        })

    rabbit_pattern = re.compile(
        r'@RabbitListener\s*\(\s*([^)]+)\)'
        r'\s*'
        + COMMON_METHOD_TAIL,
        re.MULTILINE
    )
    for m in rabbit_pattern.finditer(content):
        attrs = m.group(1)
        return_type = simple_name(m.group("return_type"))
        method_name = m.group("method_name")
        params_str = m.group("params")
        line_idx = content[:m.start()].count("\n")
        comment = extract_method_comment(lines, line_idx)
        param_types = extract_method_params(params_str)
        queue_m = re.search(r'queues\s*=\s*["\']([^"\']+)["\']', attrs)
        descriptor = f"queue={queue_m.group(1)}" if queue_m else attrs.strip()
        methods.append({
            "method_name": method_name,
            "entry_annotation": "RabbitListener",
            "entry_descriptor": descriptor,
            "return_type": return_type,
            "param_types": param_types,
            "comment": comment,
        })

    # RocketMQ @RocketMQMessageListener 是类级注解，类内 public 方法视为消息处理入口
    rocket_class_pattern = re.compile(
        r'@RocketMQMessageListener\s*\(\s*([^)]+)\)',
        re.MULTILINE
    )
    if rocket_class_pattern.search(content):
        attrs = rocket_class_pattern.search(content).group(1)
        topic_m = re.search(r'topic\s*=\s*["\']([^"\']+)["\']', attrs)
        group_m = re.search(r'consumerGroup\s*=\s*["\']([^"\']+)["\']', attrs)
        parts = []
        if topic_m:
            parts.append(f"topic={topic_m.group(1)}")
        if group_m:
            parts.append(f"consumerGroup={group_m.group(1)}")
        descriptor = ", ".join(parts) if parts else attrs.strip()

        method_pattern = re.compile(
            r'public\s+'
            r'(?:static\s+)?'
            r'(?:synchronized\s+)?'
            r'(?P<return_type>\S+(?:\s*<\s*\S+\s*(?:,\s*\S+\s*)*>)?)\s+'
            r'(?P<method_name>\w+)\s*'
            r'\((?P<params>[^)]*)\)',
            re.MULTILINE
        )
        seen = set()
        for m in method_pattern.finditer(content):
            method_name = m.group("method_name")
            if method_name in seen:
                continue
            seen.add(method_name)
            params_str = m.group("params")
            param_types = extract_method_params(params_str)
            # 只保留接收业务消息参数的方法（过滤 setter/getter）
            if not param_types:
                continue
            return_type = simple_name(m.group("return_type"))
            line_idx = content[:m.start()].count("\n")
            comment = extract_method_comment(lines, line_idx)
            methods.append({
                "method_name": method_name,
                "entry_annotation": "RocketMQMessageListener",
                "entry_descriptor": descriptor,
                "return_type": return_type,
                "param_types": param_types,
                "comment": comment,
            })

    return methods


def parse_entry_file(filepath: str, entry_type: str) -> dict | None:
    """解析单个入口文件"""
    try:
        content = Path(filepath).read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return None

    lines = content.split("\n")
    class_mapping = extract_request_mapping(content) if entry_type == "Controller" else None

    if entry_type == "Controller":
        methods = parse_controller_methods(content, lines)
    elif entry_type == "Job":
        methods = parse_job_methods(content, lines)
    elif entry_type == "Consumer":
        methods = parse_consumer_methods(content, lines)
    else:
        methods = []

    return {
        "entry_type": entry_type,
        "class_mapping": class_mapping,
        "method_count": len(methods),
        "methods": methods,
    }


def main():
    if len(sys.argv) < 2:
        print(__doc__, file=sys.stderr)
        sys.exit(1)

    project_root = os.path.abspath(sys.argv[1])
    output_path = None

    args = sys.argv[2:]
    i = 0
    while i < len(args):
        if args[i] == "--output" and i + 1 < len(args):
            output_path = args[i + 1]
            i += 2
        else:
            print(f"未知参数: {args[i]}", file=sys.stderr)
            sys.exit(1)

    if output_path is None:
        output_path = os.path.join(project_root, "knowledge", "entry-methods.json")

    entries = load_entry_files(project_root)
    type_counts = {t: 0 for t in ENTRY_FILE_TYPES.values()}
    for e in entries:
        type_counts[e["entry_type"]] += 1
    print(f"从 file-list.xlsx 召回入口: {len(entries)} 个 "
          f"(Controller={type_counts['Controller']}, "
          f"Job={type_counts['Job']}, "
          f"Consumer={type_counts['Consumer']})")

    results = []
    found = 0
    skipped = 0
    total_methods = 0

    for entry in entries:
        filepath = find_entry_path(project_root, entry["package_name"], entry["file_name"])
        if filepath is None:
            skipped += 1
            continue

        parsed = parse_entry_file(filepath, entry["entry_type"])
        if parsed is None:
            skipped += 1
            continue

        found += 1
        total_methods += parsed["method_count"]

        results.append({
            "entry_type": entry["entry_type"],
            "package_prefix": extract_package_prefix(entry["package_name"]),
            "package_name": entry["package_name"],
            "class_name": entry["file_name"],
            "class_mapping": parsed["class_mapping"],
            "method_count": parsed["method_count"],
            "methods": parsed["methods"],
        })

    output = {
        "project_root": project_root,
        "entry_count": len(entries),
        "parsed_count": found,
        "skipped_count": skipped,
        "total_methods": total_methods,
        "by_type": type_counts,
        "entries": results,
    }

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"解析完成: {found} 个入口, {total_methods} 个方法")
    print(f"跳过: {skipped} 个 (源文件未找到)")
    print(f"输出: {output_path}")


if __name__ == "__main__":
    main()

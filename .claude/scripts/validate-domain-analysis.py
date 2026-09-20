#!/usr/bin/env python3
"""
Validate s4-domain-analysis.json against S4 contracts.

用法: python3 validate-domain-analysis.py <s4_json> [--entry-list <entry-list.xlsx>]

校验规则（S4 domain-identifier agent 产出契约）:
  1. 每个候选域必须有 domain_id（英文）+ domain_name（中文），均非空
  2. 入口文件（Controller/Job/Consumer）的 responsibility_desc 必须非空
  3. 入口覆盖完整性: entry-list 的入口类 ⊆ s4 归域入口类（防漏 Consumer/Job/Controller）
  4. 提示: s4 JSON 含非入口文件归属（不报错，Excel 会过滤，S5 也会重做）

退出码: 0 通过; 1 失败; 2 输入错误
"""
import argparse
import json
import os
import sys

ENTRY_TYPES = {"Controller", "Job", "Consumer"}


def load_entry_classes(entry_list_path: str):
    """从 entry-list.xlsx 提取入口类名集合（去重），用于覆盖对账。返回 None 表示无法读取。"""
    if not entry_list_path or not os.path.exists(entry_list_path):
        return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(entry_list_path, read_only=True)
    except Exception:
        return None
    ws = wb.active
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers) if h}
    name_k = next((h for h in headers if "入口名" in str(h)), None)
    if name_k is None:
        wb.close()
        return None
    classes = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        entry_name = str(row[idx[name_k]] or "")
        parts = entry_name.split(".")
        if len(parts) >= 2:
            classes.add(parts[-2])  # package.Class.method -> Class
    wb.close()
    return classes


def main():
    ap = argparse.ArgumentParser(description="校验 s4-domain-analysis.json S4 契约")
    ap.add_argument("s4_json", help="s4-domain-analysis.json 路径")
    ap.add_argument("--entry-list", help="entry-list.xlsx 路径（用于入口覆盖对账）")
    args = ap.parse_args()

    if not os.path.exists(args.s4_json):
        print("错误: s4 JSON 不存在: %s" % args.s4_json, file=sys.stderr)
        sys.exit(2)

    with open(args.s4_json, encoding="utf-8") as f:
        data = json.load(f)

    errors = []
    cds = data.get("candidate_domains", [])

    # 1. 域命名：domain_id + domain_name 非空
    for d in cds:
        did = str(d.get("domain_id", "")).strip()
        dname = str(d.get("domain_name", "")).strip()
        if not did:
            errors.append("域缺少 domain_id: %r" % d)
        if not dname:
            errors.append("域 %r 缺少中文 domain_name" % did)

    # 2. 入口描述非空 + 收集 s4 归域入口类名
    s4_entry_classes = set()
    entry_empty_desc = []
    non_entry_count = 0
    for d in cds:
        did = d.get("domain_id", "")
        for fobj in d.get("files", []):
            ft = fobj.get("file_type", "")
            fname = str(fobj.get("file_name", ""))
            cls = fname.replace(".java", "").replace(".kt", "")
            if ft in ENTRY_TYPES:
                s4_entry_classes.add(cls)
                if not str(fobj.get("responsibility_desc", "")).strip():
                    entry_empty_desc.append("%s (%s)" % (fname, did))
            else:
                non_entry_count += 1

    if entry_empty_desc:
        errors.append("入口描述为空: %d 个（应非空）: %s" % (len(entry_empty_desc), entry_empty_desc[:5]))

    # 3. 入口覆盖对账
    entry_classes = None
    if args.entry_list:
        entry_classes = load_entry_classes(args.entry_list)
        if entry_classes is None:
            print("警告: 无法读取 entry-list，跳过覆盖对账: %s" % args.entry_list)
        else:
            missing = entry_classes - s4_entry_classes
            if missing:
                errors.append("入口覆盖缺失: entry-list 中 %d 个入口类未在 s4 归域: %s"
                              % (len(missing), sorted(missing)[:10]))

    # 4. 提示非入口（不报错）
    if non_entry_count:
        print("提示: s4 JSON 含 %d 个非入口文件归属（Entity/Mapper/Service，Excel 会过滤，S5 重做）" % non_entry_count)

    # 汇总
    print("候选域: %d 个" % len(cds))
    print("s4 归域入口类: %d 个" % len(s4_entry_classes))
    if entry_classes is not None:
        print("entry-list 入口类: %d 个" % len(entry_classes))
    print("校验失败项: %d" % len(errors))
    if errors:
        print("--- 失败明细 ---")
        for e in errors:
            print("  - %s" % e)
        sys.exit(1)
    print("✅ S4 契约校验通过（域命名/入口描述/覆盖完整性）")


if __name__ == "__main__":
    main()

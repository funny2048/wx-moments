#!/usr/bin/env python3
"""
Write Domain Analysis — 将 Domain Identifier Agent 的推理结果写入归档 JSON + 人工 Excel（S4 产出）。

用法: python3 write-domain-analysis.py <project_root> --input <agent_json> [--log-dir <path>]

输入 JSON Schema（Agent 产出）:
  {
    "candidate_domains": [
      {
        "domain_id": "campaign",
        "domain_name": "活动域",
        "confidence": 0.9,
        "evidence": ["入口文件涉及活动管理", "包路径包含 campaign"],
        "files": [
          {"file_name": "...", "file_type": "Controller", "responsibility_desc": "...", "domain": "campaign"}
        ]
      }
    ],
    "uncertainties": [
      {"file_name": "...", "reason": "...", "candidate_domains": ["a", "b"]}
    ]
  }

产出:
  - <log_dir>/s4-domain-analysis.json   完整推理归档
  - <project_root>/knowledge/domain-analysis.xlsx   人工审核用 Excel
    列: 文件名 | 文件分类 | 职责描述 | 所属域 | 校验人
"""

import json
import os
import sys
from collections import defaultdict
from glob import glob

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误: 需要安装 openpyxl — pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)


def find_latest_log_dir(project_root: str) -> str | None:
    knowledge_dir = os.path.join(project_root, "knowledge")
    if not os.path.isdir(knowledge_dir):
        return None
    candidates = [d for d in glob(os.path.join(knowledge_dir, "domain-init-log-*"))
                  if os.path.isdir(d) and os.path.basename(d) != "domain-init-log-manual"]
    if not candidates:
        return None
    return max(candidates, key=os.path.getmtime)


def load_input(path: str) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def write_archive(data: dict, log_dir: str) -> str:
    os.makedirs(log_dir, exist_ok=True)
    out = os.path.join(log_dir, "s4-domain-analysis.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return out


ENTRY_FILE_TYPES = {"Controller", "Job", "Consumer"}


def build_entry_desc_map(entry_list_path: str) -> dict:
    """从 S3 entry-list.xlsx 聚合每个入口类的方法级描述，用于补全入口文件级空描述。

    entry-list.xlsx 列：入口类型|入口名|职责描述|校验人，入口名格式 package.Class.method。
    返回 {ClassName: "desc1；desc2；..."}（去重，最多 3 条）。
    """
    if not os.path.exists(entry_list_path):
        return {}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(entry_list_path, read_only=True)
    except Exception:
        return {}
    ws = wb.active
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers) if h}
    name_k = next((h for h in headers if "入口名" in str(h)), None)
    desc_k = next((h for h in headers if "描述" in str(h) or "职责" in str(h)), None)
    if name_k is None or desc_k is None:
        wb.close()
        return {}
    bucket = defaultdict(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        entry_name = str(row[idx[name_k]] or "")
        desc = str(row[idx[desc_k]] or "").strip()
        if not desc:
            continue
        parts = entry_name.split(".")
        if len(parts) >= 2:
            bucket[parts[-2]].append(desc)  # package.Class.method -> Class
    wb.close()
    result = {}
    for cls, descs in bucket.items():
        seen = []
        for d in descs:
            if d not in seen:
                seen.append(d)
        result[cls] = "；".join(seen[:3])
    return result


def write_excel(data: dict, output_path: str) -> int:
    # S3 entry-list 方法级描述，用于补全入口文件级空描述
    entry_desc_map = build_entry_desc_map(os.path.join(os.path.dirname(output_path), "entry-list.xlsx"))

    rows = []
    skipped_non_entry = 0
    for d in data.get("candidate_domains", []):
        did = d.get("domain_id", "")
        dname = d.get("domain_name", "")
        domain_label = f"{dname}({did})" if dname and dname != did else did
        for f in d.get("files", []):
            ft = f.get("file_type", "")
            if ft not in ENTRY_FILE_TYPES:
                skipped_non_entry += 1
                continue  # S4 只展示入口文件；Entity/Mapper/Service 归属交给 S5
            fname = f.get("file_name", "")
            desc = str(f.get("responsibility_desc", "")).strip()
            if not desc:
                desc = entry_desc_map.get(fname.replace(".java", "").replace(".kt", ""), "")
            rows.append({
                "文件名": fname,
                "文件分类": ft,
                "职责描述": desc,
                "所属域": domain_label,
                "校验人": "",
            })
    # 注：uncertainties（域边界/分类决策项）不写入 Excel —— 它们是抽象决策问题，
    # 无对应真实入口文件，写成行会产出占位符文件名（如 [xxx 边界]）误导审核。
    # 完整 uncertainties 保留在 s4-domain-analysis.json 归档，由主 Agent 在对话中与用户确认。

    if skipped_non_entry:
        print(f"过滤非入口文件(Entity/Mapper/Service 等): {skipped_non_entry} 条（S4 只展示入口，实体归属交给 S5）")

    wb = Workbook()
    ws = wb.active
    ws.title = "领域分析"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["文件名", "文件分类", "职责描述", "所属域", "校验人"]
    widths = [60, 14, 50, 18, 12]
    for col_idx, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    ws.row_dimensions[1].height = 28

    rows.sort(key=lambda r: (r["所属域"], r["文件分类"], r["文件名"]))
    for row_idx, r in enumerate(rows, start=2):
        for col_idx, key in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=r[key])
            cell.border = border
            wrap = key in ("职责描述",)
            cell.alignment = Alignment(wrap_text=True, vertical="top") if wrap else Alignment(vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return len(rows)


def print_summary(data: dict, archive_path: str, excel_path: str, row_count: int):
    domains = data.get("candidate_domains", [])
    uncertainties = data.get("uncertainties", [])
    print(f"归档 JSON: {archive_path}")
    print(f"分析 Excel: {excel_path}")
    print(f"候选域: {len(domains)} 个")
    by_conf = defaultdict(int)
    for d in domains:
        c = d.get("confidence", 0)
        bucket = "高" if c >= 0.8 else ("中" if c >= 0.5 else "低")
        by_conf[bucket] += 1
    for b in ("高", "中", "低"):
        if by_conf[b]:
            print(f"  置信度 {b}: {by_conf[b]} 个")
    print(f"文件归属记录: {row_count} 条")
    print(f"不确定项: {len(uncertainties)} 条")


def main():
    if len(sys.argv) < 2:
        print(__doc__, file=sys.stderr)
        sys.exit(1)

    project_root = os.path.abspath(sys.argv[1])
    input_path = None
    log_dir = None

    args = sys.argv[2:]
    i = 0
    while i < len(args):
        if args[i] == "--input" and i + 1 < len(args):
            input_path = args[i + 1]
            i += 2
        elif args[i] == "--log-dir" and i + 1 < len(args):
            log_dir = args[i + 1]
            i += 2
        else:
            print(f"未知参数: {args[i]}", file=sys.stderr)
            sys.exit(1)

    if not input_path:
        print("错误: 必须指定 --input <agent_json>", file=sys.stderr)
        sys.exit(1)

    if not log_dir:
        log_dir = find_latest_log_dir(project_root)
    if not log_dir:
        log_dir = os.path.join(project_root, "knowledge", f"domain-init-log-manual")
        print(f"警告: 未找到 log_dir，使用 {log_dir}", file=sys.stderr)

    try:
        data = load_input(input_path)
    except FileNotFoundError as e:
        print(f"错误: {e}", file=sys.stderr)
        sys.exit(2)

    archive_path = write_archive(data, log_dir)
    excel_path = os.path.join(project_root, "knowledge", "domain-analysis.xlsx")
    row_count = write_excel(data, excel_path)
    print_summary(data, archive_path, excel_path, row_count)


if __name__ == "__main__":
    main()

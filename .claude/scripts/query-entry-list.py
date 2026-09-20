#!/usr/bin/env python3
"""
Entry List Query — 从 entry-list.xlsx 按条件查询程序入口（S4 输入）。

用法: python3 query-entry-list.py <project_root> [选项]

选项:
  --type <types>            按入口类型过滤，逗号分隔。如: Controller,Job
                            可选值: Controller,Job,Consumer
  --class <pattern>         按入口名（FQN）正则过滤。如: Campaign.*Controller
  --desc-less-than <N>      筛选职责描述字数 < N 的入口（待补充）
  --desc-empty              仅输出职责描述为空的入口
  --limit <N>               限制输出条数
  --json                    以 JSON 格式输出

示例:
  python3 query-entry-list.py . --type Controller
  python3 query-entry-list.py . --class "Campaign.*"
  python3 query-entry-list.py . --desc-empty
  python3 query-entry-list.py . --type Controller --json
"""

import json
import os
import re
import sys
from collections import defaultdict

try:
    from openpyxl import load_workbook
except ImportError:
    print("错误: 需要安装 openpyxl — pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)


def load_entry_list(project_root: str) -> list[dict]:
    path = os.path.join(project_root, "knowledge", "entry-list.xlsx")
    if not os.path.exists(path):
        raise FileNotFoundError(f"入口清单不存在: {path}（请先执行 S3 程序入口探查）")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return []
    header = [str(c).strip() if c else "" for c in rows[0]]
    records = []
    for row in rows[1:]:
        if not any(row):
            continue
        rec = dict(zip(header, [("" if c is None else str(c).strip()) for c in row]))
        records.append(rec)
    return records


def query(records, types=None, class_pattern=None, desc_less_than=None, desc_empty=False):
    results = []
    class_re = re.compile(class_pattern, re.IGNORECASE) if class_pattern else None
    for r in records:
        rtype = r.get("入口类型", "")
        rname = r.get("入口名", "")
        rdesc = r.get("职责描述", "")
        if types and rtype not in types:
            continue
        if class_re and not class_re.search(rname):
            continue
        if desc_empty and rdesc.strip():
            continue
        if desc_less_than is not None and len(rdesc) >= desc_less_than:
            continue
        results.append(r)
    return results


def print_table(records):
    by_type = defaultdict(list)
    for r in records:
        by_type[r.get("入口类型", "未知")].append(r)
    if not by_type:
        print("  (无匹配结果)")
        return
    for t in sorted(by_type):
        items = by_type[t]
        print(f"\n--- {t} ({len(items)}) ---")
        for r in sorted(items, key=lambda x: x.get("入口名", "")):
            desc = r.get("职责描述", "")
            if len(desc) > 80:
                desc = desc[:77] + "..."
            print(f"  {r.get('入口名', '')}")
            if desc:
                print(f"    └─ {desc}")


def main():
    if len(sys.argv) < 2:
        print(__doc__, file=sys.stderr)
        sys.exit(1)

    project_root = os.path.abspath(sys.argv[1])
    types = None
    class_pattern = None
    desc_less_than = None
    desc_empty = False
    limit = None
    output_json = False

    args = sys.argv[2:]
    i = 0
    while i < len(args):
        if args[i] == "--type" and i + 1 < len(args):
            types = [t.strip() for t in args[i + 1].split(",")]
            i += 2
        elif args[i] == "--class" and i + 1 < len(args):
            class_pattern = args[i + 1]
            i += 2
        elif args[i] == "--desc-less-than" and i + 1 < len(args):
            desc_less_than = int(args[i + 1])
            i += 2
        elif args[i] == "--desc-empty":
            desc_empty = True
            i += 1
        elif args[i] == "--limit" and i + 1 < len(args):
            limit = int(args[i + 1])
            i += 2
        elif args[i] == "--json":
            output_json = True
            i += 1
        else:
            print(f"未知参数: {args[i]}", file=sys.stderr)
            sys.exit(1)

    try:
        records = load_entry_list(project_root)
    except FileNotFoundError as e:
        print(f"错误: {e}", file=sys.stderr)
        sys.exit(2)

    results = query(records, types, class_pattern, desc_less_than, desc_empty)
    if limit:
        results = results[:limit]

    if output_json:
        print(json.dumps(results, ensure_ascii=False, indent=2))
    else:
        print(f"查询结果: {len(results)} 条" + (f" (限制 {limit})" if limit else ""))
        print("-" * 60)
        print_table(results)
        print("-" * 60)
        print(f"合计: {len(results)} 条")


if __name__ == "__main__":
    main()

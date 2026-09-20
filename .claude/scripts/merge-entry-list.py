#!/usr/bin/env python3
"""
Merge S3 entry-explorer JSON outputs into entry-list.xlsx.

用法:
    python3 merge-entry-list.py <project_root> [--input-dir <path>] [--output <path>]

输入:
    knowledge/domain-init-log-{date}-{time}/s3-entry-*.json
    （或用 --input-dir 显式指定）

输出:
    knowledge/entry-list.xlsx

每个 s3-entry-*.json 的 schema:
    {
      "entry_type": "Controller|Job|Consumer",
      "class_name": "...",
      "method_name": "...",
      "package": "...",
      "http_method": "POST",         # Controller 才有
      "http_path": "/api/...",        # Controller 才有
      "responsibility_desc": "...",
      "trace": {
        "root": "...",
        "depth": 4,
        "leaves": ["XxxMapper.yyy", ...],
        "tree": { "node": "...", "file": "...", "children": [...] }
      },
      "files_touched": ["...", ...]
    }
"""

import argparse
import glob
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def get_entry_name(o: dict) -> str:
    """拼接 FQN：package.class_name.method_name，过滤空段"""
    parts = [o.get("package", ""), o.get("class_name", ""), o.get("method_name", "")]
    return ".".join(p for p in parts if p)


# 列定义：(列名, 取值函数, 列宽, 是否换行)
COLUMNS = [
    ("入口类型", lambda o: o.get("entry_type", ""),         12, False),
    ("入口名",   get_entry_name,                            70, False),
    ("职责描述", lambda o: o.get("responsibility_desc", ""), 70, True),
    ("校验人",   lambda o: "",                              12, False),
]


def find_latest_log_dir(project_root: str) -> str | None:
    """自动找最新的 knowledge/domain-init-log-YYYYMMDD-HHMMSS 目录。

    按 mtime 选最新（不依赖目录名字典序），并排除 domain-init-log-manual 兜底目录。
    """
    knowledge_dir = os.path.join(project_root, "knowledge")
    if not os.path.isdir(knowledge_dir):
        return None
    candidates = [d for d in glob.glob(os.path.join(knowledge_dir, "domain-init-log-*"))
                  if os.path.isdir(d) and os.path.basename(d) != "domain-init-log-manual"]
    if not candidates:
        return None
    return max(candidates, key=os.path.getmtime)


def load_entry_jsons(input_dir: str) -> tuple[list[dict], list[str]]:
    """读取目录下所有 s3-entry-*.json，返回 (成功列表, 失败信息列表)"""
    pattern = os.path.join(input_dir, "s3-entry-*.json")
    files = sorted(glob.glob(pattern))

    entries: list[dict] = []
    failed: list[str] = []
    for f in files:
        try:
            with open(f, "r", encoding="utf-8") as fp:
                data = json.load(fp)
            if not isinstance(data, dict):
                raise ValueError("根元素必须是 JSON 对象")
            entries.append(data)
        except Exception as e:
            failed.append(f"{os.path.basename(f)}: {e}")
    return entries, failed


def write_excel(entries: list[dict], output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "入口清单"

    # 样式
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 表头
    for col_idx, (col_name, _, _, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[1].height = 28

    # 数据行
    for row_idx, entry in enumerate(entries, start=2):
        for col_idx, (_, getter, _, wrap) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=getter(entry))
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top") if wrap else Alignment(vertical="center")

    # 列宽
    for col_idx, (_, _, width, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 冻结首行 + 自动筛选
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)


def print_summary(entries: list[dict], failed: list[str], input_dir: str, output_path: str) -> None:
    by_type: dict[str, int] = {}
    for e in entries:
        t = e.get("entry_type", "未知")
        by_type[t] = by_type.get(t, 0) + 1

    empty_desc  = sum(1 for e in entries if not str(e.get("responsibility_desc", "")).strip())
    long_desc   = sum(1 for e in entries if len(str(e.get("responsibility_desc", ""))) > 100)
    empty_trace = sum(1 for e in entries if not (e.get("trace") or {}).get("tree"))
    empty_files = sum(1 for e in entries if not (e.get("files_touched") or []))

    print(f"输入: {input_dir}")
    print(f"读取: {len(entries)} 个入口 JSON")
    for t in sorted(by_type):
        print(f"  {t}: {by_type[t]}")
    if failed:
        print(f"警告: {len(failed)} 个文件解析失败:")
        for msg in failed:
            print(f"  - {msg}")
    print(f"输出: {output_path}")
    print("健康度:")
    print(f"  职责描述为空:    {empty_desc}")
    print(f"  职责描述>100字:  {long_desc}")
    print(f"  trace.tree 为空: {empty_trace}")
    print(f"  files_touched 空: {empty_files}")


def main() -> None:
    parser = argparse.ArgumentParser(description="合并 S3 entry JSON 到 entry-list.xlsx")
    parser.add_argument("project_root", help="目标项目根目录")
    parser.add_argument("--input-dir", help="s3-entry-*.json 所在目录（默认自动找最新 log 目录）")
    parser.add_argument("--output", help="输出 xlsx 路径（默认: knowledge/entry-list.xlsx）")
    args = parser.parse_args()

    project_root = os.path.abspath(args.project_root)

    if args.input_dir:
        input_dir = args.input_dir
    else:
        input_dir = find_latest_log_dir(project_root)
        if input_dir is None:
            print(f"错误: 未找到 {project_root}/knowledge/domain-init-log-* 目录", file=sys.stderr)
            print("提示: 用 --input-dir 显式指定 JSON 所在目录", file=sys.stderr)
            sys.exit(1)

    if not os.path.isdir(input_dir):
        print(f"错误: 输入目录不存在: {input_dir}", file=sys.stderr)
        sys.exit(1)

    output_path = args.output or os.path.join(project_root, "knowledge", "entry-list.xlsx")

    entries, failed = load_entry_jsons(input_dir)
    if not entries:
        print(f"错误: {input_dir} 下未找到有效的 s3-entry-*.json", file=sys.stderr)
        sys.exit(1)

    # 排序：入口类型 → 类名 → 方法名
    entries.sort(key=lambda e: (
        e.get("entry_type", ""),
        e.get("class_name", ""),
        e.get("method_name", ""),
    ))

    write_excel(entries, output_path)
    print_summary(entries, failed, input_dir, output_path)


if __name__ == "__main__":
    main()
